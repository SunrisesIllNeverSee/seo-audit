#!/usr/bin/env python3
"""
generate_projections.py — read the SEO audit workbook and emit
agent-readable YAML + Markdown projections.

Usage:
    python3 scripts/generate_projections.py

Source: source/ello-cello-seo-audit.xlsx
Output: projections/*.yaml, reports/*.md
"""
import sys
import os
from pathlib import Path
from datetime import date

try:
    import openpyxl
except ImportError:
    print("openpyxl is required: pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "source" / "ello-cello-seo-audit.xlsx"
PROJECTIONS = ROOT / "projections"
REPORTS = ROOT / "reports"


def cell(val):
    """Clean a cell value for YAML/Markdown output."""
    if val is None:
        return ""
    return str(val).strip()


def sheet_rows(ws, header_row=1):
    """Yield dicts from a worksheet, using header_row as keys."""
    headers = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        headers = [cell(c) for c in row]
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(cell(c) for c in row):
            continue
        yield {headers[i]: cell(row[i]) for i in range(len(headers)) if i < len(row)}


def to_yaml_scalar(val):
    """Convert a string to a safe YAML scalar."""
    s = str(val)
    if not s:
        return '""'
    # Quote if it contains special chars
    if any(c in s for c in [":", "#", "'", '"', "\n", "{", "}", "[", "]", ",", "&", "*", "?", "|", ">", "@", "`"]):
        return '"' + s.replace('"', '\\"') + '"'
    return s


def yaml_list(items, indent=2):
    """Render a list of dicts as YAML."""
    lines = []
    pad = " " * indent
    for item in items:
        lines.append(f"{pad}-")
        for key, val in item.items():
            lines.append(f"{pad}  {key}: {to_yaml_scalar(val)}")
    return "\n".join(lines)


def yaml_kv(key, val, indent=0):
    """Render a key-value pair as YAML."""
    pad = " " * indent
    return f"{pad}{key}: {to_yaml_scalar(val)}"


# ─── Sheet extractors ────────────────────────────────────────────────

def extract_read_me(wb):
    ws = wb["Read Me"]
    items = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] and cell(row[0]) != "Item":
            items[cell(row[0])] = cell(row[1])
    return items


def extract_actions(wb):
    ws = wb["Actions"]
    return list(sheet_rows(ws))


def extract_competitors(wb):
    ws = wb["Competitors"]
    return list(sheet_rows(ws))


def extract_keyword_map(wb):
    ws = wb["Keyword Map"]
    return list(sheet_rows(ws))


def extract_page_drafts(wb):
    ws = wb["Page Drafts"]
    return list(sheet_rows(ws))


def extract_consolidation(wb):
    ws = wb["Consolidation"]
    return list(sheet_rows(ws))


def extract_domain_inventory(wb):
    ws = wb["Domain Inventory"]
    return list(sheet_rows(ws))


def extract_page_metrics(wb):
    ws = wb["Page Metrics"]
    return list(sheet_rows(ws))


def extract_url_checks(wb):
    ws = wb["URL Checks"]
    return list(sheet_rows(ws))


# ─── YAML writers ────────────────────────────────────────────────────

def write_yaml(filename, content):
    path = PROJECTIONS / filename
    path.write_text(content)
    print(f"  projections/{filename}")


def emit_read_me(read_me):
    lines = ["# Read Me — audit metadata"]
    lines.append("")
    for key, val in read_me.items():
        lines.append(yaml_kv(key, val))
    write_yaml("read-me.yaml", "\n".join(lines) + "\n")


def emit_actions(actions):
    lines = ["# Actions — prioritized SEO recommendations"]
    lines.append("")
    lines.append("actions:")
    lines.append(yaml_list(actions))
    write_yaml("actions.yaml", "\n".join(lines) + "\n")


def emit_competitors(competitors):
    lines = ["# Competitors — verified competitive landscape"]
    lines.append("")
    lines.append("competitors:")
    lines.append(yaml_list(competitors))
    write_yaml("competitors.yaml", "\n".join(lines) + "\n")


def emit_keyword_map(keyword_map):
    lines = ["# Keyword Map — query clusters to preferred URLs"]
    lines.append("")
    lines.append("keyword_map:")
    lines.append(yaml_list(keyword_map))
    write_yaml("keyword-map.yaml", "\n".join(lines) + "\n")


def emit_page_drafts(page_drafts):
    lines = ["# Page Drafts — proposed title/description/heading rewrites"]
    lines.append("")
    lines.append("page_drafts:")
    lines.append(yaml_list(page_drafts))
    write_yaml("page-drafts.yaml", "\n".join(lines) + "\n")


def emit_consolidation(consolidation):
    lines = ["# Consolidation — redirect/merge candidates"]
    lines.append("")
    lines.append("consolidation:")
    lines.append(yaml_list(consolidation))
    write_yaml("consolidation.yaml", "\n".join(lines) + "\n")


def emit_domain_inventory(domains):
    lines = ["# Domain Inventory — domains in scope with roles"]
    lines.append("")
    lines.append("domains:")
    lines.append(yaml_list(domains))
    write_yaml("domain-inventory.yaml", "\n".join(lines) + "\n")


def emit_page_metrics(metrics):
    lines = ["# Page Metrics — GSC data for key pages"]
    lines.append("")
    lines.append("page_metrics:")
    lines.append(yaml_list(metrics))
    write_yaml("page-metrics.yaml", "\n".join(lines) + "\n")


def emit_url_checks(url_checks):
    lines = ["# URL Checks — 76 URL inspection results"]
    lines.append("")
    lines.append("url_checks:")
    lines.append(yaml_list(url_checks))
    write_yaml("url-checks.yaml", "\n".join(lines) + "\n")


# ─── Markdown report writers ─────────────────────────────────────────

def write_md(filename, content):
    path = REPORTS / filename
    path.write_text(content)
    print(f"  reports/{filename}")


def emit_overview(read_me, actions, domains, metrics):
    lines = [
        "# SEO Audit — Executive Overview",
        "",
        f"**Audit date:** {read_me.get('Audit date', 'unknown')}",
        "",
        f"**Scope:** {read_me.get('Scope', '')}",
        "",
        "## Domains in scope",
        "",
        "| Domain | Role | Sitemap URLs |",
        "|--------|------|-------------|",
    ]
    for d in domains:
        lines.append(f"| {d.get('Domain','')} | {d.get('Recommended role','')} | {d.get('Sitemap URLs','')} |")
    lines.append("")
    lines.append("## Top priorities")
    lines.append("")
    lines.append("| # | Action | Impact | Effort |")
    lines.append("|---|--------|--------|--------|")
    for a in actions:
        lines.append(f"| P{a.get('Priority','')} | {a.get('Action','')} | {a.get('Impact','')} | {a.get('Work estimate','')} |")
    lines.append("")
    lines.append("## Key page metrics (last 90 days)")
    lines.append("")
    lines.append("| Page | Impressions | Clicks | CTR |")
    lines.append("|------|-------------|--------|-----|")
    for m in metrics:
        ctr = m.get("CTR", "")
        if ctr:
            try:
                ctr = f"{float(ctr)*100:.1f}%"
            except ValueError:
                pass
        lines.append(f"| {m.get('Page','')} | {m.get('Impressions','')} | {m.get('Clicks','')} | {ctr} |")
    lines.append("")
    lines.append("## Evidence labels")
    lines.append("")
    lines.append(f"- **Verified** — {read_me.get('Evidence labels', '')}")
    lines.append(f"- **Limits** — {read_me.get('Limits', '')}")
    lines.append(f"- **Historical baseline** — {read_me.get('Historical baseline', '')}")
    lines.append("")
    write_md("overview.md", "\n".join(lines))


def emit_actions_report(actions):
    lines = ["# SEO Audit — Action Plan", ""]
    for a in actions:
        lines.append(f"## P{a.get('Priority','')} — {a.get('Action','')}")
        lines.append("")
        lines.append(f"**Impact:** {a.get('Impact','')}")
        lines.append(f"**Effort:** {a.get('Work estimate','')}")
        lines.append(f"**Estimated results:** {a.get('Estimated results','')}")
        lines.append(f"**Confidence:** {a.get('Confidence','')}")
        lines.append("")
        lines.append(f"**Affected pages:** {a.get('Affected pages','')}")
        lines.append("")
        lines.append(f"**Evidence:** {a.get('Evidence','')}")
        lines.append("")
        lines.append(f"**Recommended change:** {a.get('Recommended change','')}")
        lines.append("")
        lines.append(f"**Acceptance check:** {a.get('Acceptance check','')}")
        lines.append("")
        lines.append("---")
        lines.append("")
    write_md("actions.md", "\n".join(lines))


def emit_keyword_report(keyword_map, page_drafts):
    lines = ["# SEO Audit — Keyword Strategy", ""]
    lines.append("## Query clusters")
    lines.append("")
    lines.append("| Query cluster | Intent | Preferred URL | Impact |")
    lines.append("|---------------|--------|---------------|--------|")
    for k in keyword_map:
        lines.append(f"| {k.get('Query cluster','')} | {k.get('Search intent','')} | {k.get('Preferred URL','')} | {k.get('Impact','')} |")
    lines.append("")
    lines.append("## Proposed page rewrites")
    lines.append("")
    for p in page_drafts:
        lines.append(f"### {p.get('Page / placement','')}")
        lines.append("")
        lines.append(f"- **Current title:** {p.get('Current title / state','')}")
        lines.append(f"- **Proposed title:** {p.get('Proposed title / CTA','')}")
        lines.append(f"- **Proposed heading:** {p.get('Proposed heading','')}")
        lines.append(f"- **Proposed description:** {p.get('Proposed description / copy','')}")
        lines.append(f"- **Implementation note:** {p.get('Implementation note','')}")
        lines.append(f"- **Impact:** {p.get('Impact','')} | **ETA:** {p.get('Estimated results','')}")
        lines.append("")
    write_md("keyword-strategy.md", "\n".join(lines))


# ─── Main ────────────────────────────────────────────────────────────

def main():
    if not SOURCE.exists():
        print(f"Source workbook not found: {SOURCE}", file=sys.stderr)
        sys.exit(1)

    print(f"Reading {SOURCE.name}...")
    wb = openpyxl.load_workbook(SOURCE, data_only=True)

    read_me = extract_read_me(wb)
    actions = extract_actions(wb)
    competitors = extract_competitors(wb)
    keyword_map = extract_keyword_map(wb)
    page_drafts = extract_page_drafts(wb)
    consolidation = extract_consolidation(wb)
    domains = extract_domain_inventory(wb)
    metrics = extract_page_metrics(wb)
    url_checks = extract_url_checks(wb)

    print(f"  {len(actions)} actions, {len(competitors)} competitors, "
          f"{len(keyword_map)} keyword clusters, {len(page_drafts)} page drafts, "
          f"{len(consolidation)} consolidation candidates, {len(domains)} domains, "
          f"{len(metrics)} page metrics, {len(url_checks)} URL checks")

    print("\nGenerating YAML projections...")
    emit_read_me(read_me)
    emit_actions(actions)
    emit_competitors(competitors)
    emit_keyword_map(keyword_map)
    emit_page_drafts(page_drafts)
    emit_consolidation(consolidation)
    emit_domain_inventory(domains)
    emit_page_metrics(metrics)
    emit_url_checks(url_checks)

    print("\nGenerating Markdown reports...")
    emit_overview(read_me, actions, domains, metrics)
    emit_actions_report(actions)
    emit_keyword_report(keyword_map, page_drafts)

    print(f"\nDone. Projections in {PROJECTIONS}, reports in {REPORTS}")


if __name__ == "__main__":
    main()
